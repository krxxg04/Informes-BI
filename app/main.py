from flask import Flask, render_template, request, send_file, jsonify
from datetime import datetime
import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import json

app = Flask(__name__, 
            template_folder='../templates',
            static_folder='../static')

# Configuración para producción
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max

# Configuración de carpetas
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCUMENTOS_FOLDER = os.path.join(BASE_DIR, 'documentos_generados')
PLANTILLAS_FOLDER = os.path.join(BASE_DIR, 'plantillas_word')

# Asegurarse de que las carpetas existan
os.makedirs(DOCUMENTOS_FOLDER, exist_ok=True)
os.makedirs(PLANTILLAS_FOLDER, exist_ok=True)

@app.route('/')
def index():
    """Página principal con selector de tipo de informe"""
    return render_template('index.html')

@app.route('/health')
def health():
    """Health check endpoint para keep-alive"""
    return jsonify({
        'status': 'healthy',
        'timestamp': datetime.now().isoformat(),
        'service': 'Generador de Informes'
    }), 200

@app.route('/ping')
def ping():
    """Ping endpoint simple"""
    return 'pong', 200

@app.route('/formulario/<tipo_informe>')
def formulario(tipo_informe):
    """Muestra el formulario según el tipo de informe"""
    tipos_validos = ['expediente', 'resolucion_directoral', 'acta_conei', 'conformidad']
    
    if tipo_informe not in tipos_validos:
        return "Tipo de informe no válido", 404
    
    return render_template(f'formulario_{tipo_informe}.html', tipo=tipo_informe)

@app.route('/generar/<tipo_informe>', methods=['POST'])
def generar_documento(tipo_informe):
    """Genera el documento Word según el tipo de informe"""
    try:
        datos = request.json
        
        # Crear documento Word
        doc = Document()
        
        # Configurar márgenes
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
        
        # Generar según tipo
        if tipo_informe == 'expediente':
            generar_expediente(doc, datos)
        elif tipo_informe == 'resolucion_directoral':
            generar_resolucion_directoral(doc, datos)
        elif tipo_informe == 'acta_conei':
            generar_acta_conei(doc, datos)
        elif tipo_informe == 'conformidad':
            generar_conformidad(doc, datos)
        
        # Guardar documento
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'{tipo_informe}_{timestamp}.docx'
        filepath = os.path.join(DOCUMENTOS_FOLDER, filename)
        doc.save(filepath)
        
        return jsonify({
            'success': True,
            'filename': filename,
            'download_url': f'/descargar/{filename}'
        })
    
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500

@app.route('/descargar/<filename>')
def descargar(filename):
    """Descarga el documento generado"""
    filepath = os.path.join(DOCUMENTOS_FOLDER, filename)
    if os.path.exists(filepath):
        return send_file(filepath, as_attachment=True)
    return "Archivo no encontrado", 404

def generar_expediente(doc, datos):
    """Genera un expediente administrativo"""
    # Encabezado centrado
    header = doc.add_paragraph()
    header.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = header.add_run('EXPEDIENTE ADMINISTRATIVO\n')
    run.bold = True
    run.font.size = Pt(14)
    
    # Número de expediente
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run(f"N° {datos.get('numero_expediente', '')}")
    run.font.size = Pt(12)
    
    doc.add_paragraph()  # Espacio
    
    # Datos del expediente
    doc.add_paragraph(f"FECHA: {datos.get('fecha', datetime.now().strftime('%d/%m/%Y'))}")
    doc.add_paragraph(f"DE: {datos.get('remitente', '')}")
    doc.add_paragraph(f"PARA: {datos.get('destinatario', '')}")
    doc.add_paragraph(f"ASUNTO: {datos.get('asunto', '')}")
    
    doc.add_paragraph()  # Espacio
    
    # Contenido
    if datos.get('contenido'):
        doc.add_paragraph(datos['contenido'])
    
    # Antecedentes si existen
    if datos.get('antecedentes'):
        doc.add_paragraph()
        p = doc.add_paragraph()
        run = p.add_run('ANTECEDENTES:')
        run.bold = True
        doc.add_paragraph(datos['antecedentes'])
    
    # Análisis si existe
    if datos.get('analisis'):
        doc.add_paragraph()
        p = doc.add_paragraph()
        run = p.add_run('ANÁLISIS:')
        run.bold = True
        doc.add_paragraph(datos['analisis'])
    
    # Conclusión si existe
    if datos.get('conclusion'):
        doc.add_paragraph()
        p = doc.add_paragraph()
        run = p.add_run('CONCLUSIÓN:')
        run.bold = True
        doc.add_paragraph(datos['conclusion'])

def generar_resolucion_directoral(doc, datos):
    """Genera una Resolución Directoral"""
    # Encabezado
    header = doc.add_paragraph()
    header.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = header.add_run('RESOLUCIÓN DIRECTORAL\n')
    run.bold = True
    run.font.size = Pt(14)
    
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run(f"N° {datos.get('numero_rd', '')}")
    run.font.size = Pt(12)
    
    doc.add_paragraph()
    
    # Institución y lugar
    doc.add_paragraph(f"{datos.get('institucion', '')}")
    doc.add_paragraph(f"{datos.get('lugar', '')}, {datos.get('fecha', datetime.now().strftime('%d de %B de %Y'))}")
    
    doc.add_paragraph()
    
    # Visto
    p = doc.add_paragraph()
    run = p.add_run('VISTO:')
    run.bold = True
    doc.add_paragraph(datos.get('visto', ''))
    
    doc.add_paragraph()
    
    # Considerando
    p = doc.add_paragraph()
    run = p.add_run('CONSIDERANDO:')
    run.bold = True
    
    considerandos = datos.get('considerandos', [])
    for i, considerando in enumerate(considerandos, 1):
        doc.add_paragraph(f"Que, {considerando}")
    
    doc.add_paragraph()
    
    # SE RESUELVE
    p = doc.add_paragraph()
    run = p.add_run('SE RESUELVE:')
    run.bold = True
    
    articulos = datos.get('articulos', [])
    for i, articulo in enumerate(articulos, 1):
        doc.add_paragraph(f"Artículo {i}°.- {articulo}")
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Firma
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph('_________________________________')
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p = doc.add_paragraph(datos.get('firmante', ''))
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p = doc.add_paragraph(datos.get('cargo_firmante', ''))
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

def generar_acta_conei(doc, datos):
    """Genera un Acta de CONEI"""
    # Encabezado
    header = doc.add_paragraph()
    header.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = header.add_run('ACTA DE REUNIÓN DEL CONSEJO EDUCATIVO INSTITUCIONAL\n')
    run.bold = True
    run.font.size = Pt(14)
    
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run(f"ACTA N° {datos.get('numero_acta', '')}")
    run.font.size = Pt(12)
    
    doc.add_paragraph()
    
    # Datos básicos
    doc.add_paragraph(f"INSTITUCIÓN EDUCATIVA: {datos.get('institucion', '')}")
    doc.add_paragraph(f"FECHA: {datos.get('fecha', datetime.now().strftime('%d/%m/%Y'))}")
    doc.add_paragraph(f"HORA DE INICIO: {datos.get('hora_inicio', '')}")
    doc.add_paragraph(f"HORA DE TÉRMINO: {datos.get('hora_termino', '')}")
    doc.add_paragraph(f"LUGAR: {datos.get('lugar', '')}")
    
    doc.add_paragraph()
    
    # Asistentes
    p = doc.add_paragraph()
    run = p.add_run('ASISTENTES:')
    run.bold = True
    
    asistentes = datos.get('asistentes', [])
    for asistente in asistentes:
        doc.add_paragraph(f"• {asistente}", style='List Bullet')
    
    doc.add_paragraph()
    
    # Agenda
    p = doc.add_paragraph()
    run = p.add_run('AGENDA:')
    run.bold = True
    
    agenda = datos.get('agenda', [])
    for i, punto in enumerate(agenda, 1):
        doc.add_paragraph(f"{i}. {punto}")
    
    doc.add_paragraph()
    
    # Desarrollo
    p = doc.add_paragraph()
    run = p.add_run('DESARROLLO:')
    run.bold = True
    doc.add_paragraph(datos.get('desarrollo', ''))
    
    doc.add_paragraph()
    
    # Acuerdos
    p = doc.add_paragraph()
    run = p.add_run('ACUERDOS:')
    run.bold = True
    
    acuerdos = datos.get('acuerdos', [])
    for i, acuerdo in enumerate(acuerdos, 1):
        doc.add_paragraph(f"{i}. {acuerdo}")
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Firmas
    doc.add_paragraph('Sin otro particular que tratar, se da por concluida la reunión.')
    doc.add_paragraph()
    doc.add_paragraph()
    
    p = doc.add_paragraph('_________________________________')
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p = doc.add_paragraph('PRESIDENTE DEL CONEI')
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

def generar_conformidad(doc, datos):
    """Genera un documento de conformidad"""
    # Encabezado
    header = doc.add_paragraph()
    header.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = header.add_run('CONSTANCIA DE CONFORMIDAD\n')
    run.bold = True
    run.font.size = Pt(14)
    
    doc.add_paragraph()
    
    # Datos
    doc.add_paragraph(f"INSTITUCIÓN: {datos.get('institucion', '')}")
    doc.add_paragraph(f"FECHA: {datos.get('fecha', datetime.now().strftime('%d/%m/%Y'))}")
    
    doc.add_paragraph()
    
    # Contenido
    doc.add_paragraph(f"Yo, {datos.get('nombre_responsable', '')}, en mi calidad de {datos.get('cargo', '')}, doy conformidad a:")
    
    doc.add_paragraph()
    doc.add_paragraph(datos.get('descripcion', ''))
    
    doc.add_paragraph()
    
    # Observaciones si existen
    if datos.get('observaciones'):
        p = doc.add_paragraph()
        run = p.add_run('OBSERVACIONES:')
        run.bold = True
        doc.add_paragraph(datos['observaciones'])
    
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Firma
    p = doc.add_paragraph('_________________________________')
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p = doc.add_paragraph(datos.get('nombre_responsable', ''))
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p = doc.add_paragraph(datos.get('cargo', ''))
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

if __name__ == '__main__':
    # Configuración según el entorno
    port = int(os.environ.get('PORT', 5000))
    debug = os.environ.get('FLASK_ENV') != 'production'
    
    app.run(host='0.0.0.0', port=port, debug=debug)
